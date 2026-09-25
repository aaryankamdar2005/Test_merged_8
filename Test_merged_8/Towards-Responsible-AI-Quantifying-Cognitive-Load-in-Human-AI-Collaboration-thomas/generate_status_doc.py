import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()

    # Title
    title = doc.add_heading('Project Status Documentation', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Quantifying Cognitive Load in Human-AI Collaboration\n').alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        "This project aims to quantify and predict cognitive load in AI interfaces. "
        "Because public datasets for this specific human-AI interaction did not exist, a custom "
        "interaction logger and simulated environment were built and deployed. This logger "
        "tracks real-time user metrics (facial expressions, eye tracking, mouse movements, keystrokes) "
        "while users complete tasks. The collected data was augmented with synthetic values to create "
        "a robust dataset, which was then used to train a machine learning pipeline capable of predicting "
        "cognitive load levels."
    )

    # 2. Architecture & Tech Stack
    doc.add_heading('2. Architecture & Tech Stack', level=1)
    tech_stack = doc.add_paragraph()
    tech_stack.add_run('Frontend: ').bold = True
    tech_stack.add_run('Vite, Vanilla JavaScript (Deployed on Vercel)\n')
    tech_stack.add_run('Backend: ').bold = True
    tech_stack.add_run('FastAPI, Python (Deployed on Render)\n')
    tech_stack.add_run('Database: ').bold = True
    tech_stack.add_run('PostgreSQL via Supabase\n')
    tech_stack.add_run('ML / Data Processing: ').bold = True
    tech_stack.add_run('XGBoost, LightGBM, Scikit-learn, OpenCV, MediaPipe')

    # 3. Interaction Logger & Data Collection
    doc.add_heading('3. Interaction Logger (Simulated Environment)', level=1)
    doc.add_paragraph(
        "To collect data, a web-based simulated environment was developed. This environment acts as the "
        "interaction logger, presenting users with tasks and a chat interface while monitoring their behavior in the background."
    )
    doc.add_paragraph('Key metrics collected in real-time include:', style='List Bullet')
    doc.add_paragraph('Facial expressions (happy, neutral, frustrated) via MediaPipe/OpenCV', style='List Bullet')
    doc.add_paragraph('Eye tracking (EAR, PERCLOS, gaze direction)', style='List Bullet')
    doc.add_paragraph('Keyboard interactions (keypresses, backspaces, backspace rate, thinking pauses)', style='List Bullet')
    doc.add_paragraph('Mouse tracking (mouse moves, total cursor distance)', style='List Bullet')
    doc.add_paragraph('Interaction metrics (prompts sent, average prompt length)', style='List Bullet')

    # 4. Data Pipeline & Dataset Generation
    doc.add_heading('4. Dataset Generation & Synthetic Data', level=1)
    doc.add_paragraph(
        "The interaction logs were aggregated per participant and per task to form the foundational dataset. "
        "To enhance the robustness of the ML models and handle the lack of an initial public dataset, "
        "synthetic data generation techniques were applied. This allowed the creation of the 'Final_Dataset.csv' "
        "(12,000+ rows) with varied scenarios representing different cognitive load profiles."
    )
    doc.add_paragraph('The dataset includes the calculated Cognitive Load Index (CLI_score) and its corresponding categorical label (CLI_label):', style='List Bullet')
    doc.add_paragraph('Low', style='List Bullet 2')
    doc.add_paragraph('Moderate', style='List Bullet 2')
    doc.add_paragraph('High', style='List Bullet 2')
    doc.add_paragraph('Very High / Overload', style='List Bullet 2')

    # 5. Machine Learning Pipeline
    doc.add_heading('5. Machine Learning Pipeline', level=1)
    doc.add_paragraph(
        "The machine learning pipeline is the core crux of the project, responsible for predicting the "
        "continuous cognitive load score (CLI_score) and categorizing the user's state (CLI_label). "
        "Due to the highly dimensional and non-linear nature of human interaction data (eye tracking, facial "
        "expressions, keystrokes, and mouse movements), advanced gradient boosting and ensemble algorithms were employed."
    )
    doc.add_paragraph(
        "A grouped held-out validation strategy (GroupShuffleSplit) was used to ensure that data from the same participant "
        "did not leak across the training and testing sets. Out of the 12,000 generated samples, 2,397 were used for the "
        "held-out test set."
    )
    
    doc.add_paragraph('Model Evaluation & Accuracies:', style='List Bullet')
    doc.add_paragraph('XGBoost Classifier: 66.00% accuracy (Best Performing)', style='List Bullet 2')
    doc.add_paragraph('LightGBM Classifier: 65.87% accuracy', style='List Bullet 2')
    doc.add_paragraph('Random Forest Classifier: 65.21% accuracy', style='List Bullet 2')
    doc.add_paragraph('Extra Trees Classifier: 63.70% accuracy', style='List Bullet 2')
    
    doc.add_paragraph(
        "Feature Importance Analysis:", style='List Bullet'
    )
    doc.add_paragraph(
        "To ensure the models were interpretable, a feature importance analysis was conducted. "
        "The analysis revealed that behavioral metrics such as 'total cursor distance px', 'average prompt length words', "
        "and 'mean thinking pause seconds' heavily influence the cognitive load prediction, more so than simple "
        "facial expressions in isolation. This validates the multi-modal approach of the interaction logger."
    )

    # 6. Deployment Status
    doc.add_heading('6. Current Deployment Status', level=1)
    doc.add_paragraph(
        "The system has been successfully deployed to production environments to allow widespread access "
        "and data collection."
    )
    doc.add_paragraph('Backend (Render):', style='List Bullet')
    doc.add_paragraph('The FastAPI backend is deployed on Render.', style='List Bullet 2')
    doc.add_paragraph('Connected to a Supabase PostgreSQL database using the IPv4 Session Pooler.', style='List Bullet 2')
    doc.add_paragraph('Dependencies such as psycopg[binary], xgboost, and mediapipe have been configured to support the heavy ML processing within the deployment constraints.', style='List Bullet 2')
    
    doc.add_paragraph('Frontend (Vercel):', style='List Bullet')
    doc.add_paragraph('The Vite-based frontend is deployed on Vercel as a static site.', style='List Bullet 2')
    doc.add_paragraph('Configured to communicate with the Render backend API.', style='List Bullet 2')

    # Save Document
    file_path = os.path.join(os.getcwd(), 'Project_Status_Documentation_v2.docx')
    doc.save(file_path)
    print(f"Document saved to: {file_path}")

if __name__ == "__main__":
    create_document()
